"""
tool_layer/document_tools.py
Multi-format template reading and output writing.

Supported formats:
  - Text: .md, .txt, .sql, .json, .xml, .yaml, .csv, .html
  - Office: .docx (Word), .xlsx (Excel), .pptx (PowerPoint)

Office libraries (python-docx, openpyxl, python-pptx) should be placed in
the vendor/ directory by the user for zero-install deployment.
"""

import re
import os
from pathlib import Path

# ---------------------------------------------------------------------------
# Optional office library imports
# ---------------------------------------------------------------------------

_docx = None
_openpyxl = None
_pptx = None

def _try_imports(vendor_dir: Path):
    """Try to import office libraries from vendor/ directory."""
    global _docx, _openpyxl, _pptx

    vendor_str = str(vendor_dir)
    if vendor_str not in __import__("sys").path:
        __import__("sys").path.insert(0, vendor_str)

    # python-docx
    try:
        import docx as _d
        _docx = _d
    except ImportError:
        pass

    # openpyxl
    try:
        import openpyxl as _o
        _openpyxl = _o
    except ImportError:
        pass

    # python-pptx
    try:
        import pptx as _p
        _pptx = _p
    except ImportError:
        pass


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

_PLACEHOLDER_RE = re.compile(r"\{\{(.+?)\}\}")

_TEXT_EXTS = {".md", ".txt", ".sql", ".json", ".xml", ".yaml", ".yml", ".csv", ".html", ".htm", ".py", ".js", ".ts", ".java"}


def _extract_placeholders(text: str) -> list[str]:
    """Extract unique {{KEY}} placeholders from text."""
    return list(dict.fromkeys(m.strip() for m in _PLACEHOLDER_RE.findall(text)))


def _replace_placeholders(text: str, fields: dict) -> str:
    """Replace {{KEY}} placeholders with values from fields dict.

    Missing keys are replaced with empty string.
    """
    def replacer(m):
        key = m.group(1).strip()
        val = fields.get(key, "")
        return str(val)
    return _PLACEHOLDER_RE.sub(replacer, text)


def _ensure_output_dir(path: Path):
    """Create parent directories for output path."""
    path.parent.mkdir(parents=True, exist_ok=True)


# ---------------------------------------------------------------------------
# read_template — multi-format
# ---------------------------------------------------------------------------

def read_template(path: str, base_dir: Path) -> str:
    """Read a template file from templates/ directory.

    For text-based formats, returns the raw content with placeholder list.
    For Office formats, extracts text content and placeholders.
    Returns a formatted string describing the template.
    """
    p = Path(path)
    if not p.is_absolute():
        p = base_dir / "templates" / p

    if not p.exists():
        return "[Error] Template not found: {}".format(p)

    ext = p.suffix.lower()

    # Text-based templates
    if ext in _TEXT_EXTS:
        try:
            content = p.read_text(encoding="utf-8", errors="replace")
            placeholders = _extract_placeholders(content)
            lines = [
                "[Template] {}".format(p.name),
                "Format: {}".format(ext),
                "Size: {} chars".format(len(content)),
            ]
            if placeholders:
                lines.append("Placeholders: {}".format(", ".join("{{{{{}}}}}".format(k) for k in placeholders)))
            lines.append("\n--- Template Content ---\n")
            lines.append(content)
            return "\n".join(lines)
        except Exception as e:
            return "[Error] Cannot read template {}: {}".format(p, e)

    # .docx
    if ext == ".docx":
        return _read_docx_template(p, base_dir)

    # .xlsx
    if ext == ".xlsx":
        return _read_xlsx_template(p, base_dir)

    # .pptx
    if ext == ".pptx":
        return _read_pptx_template(p, base_dir)

    # Fallback: try as text
    try:
        content = p.read_text(encoding="utf-8", errors="replace")
        placeholders = _extract_placeholders(content)
        lines = [
            "[Template] {}".format(p.name),
            "Format: {}".format(ext),
            "Size: {} chars".format(len(content)),
        ]
        if placeholders:
            lines.append("Placeholders: {}".format(", ".join("{{{{{}}}}}".format(k) for k in placeholders)))
        lines.append("\n--- Template Content ---\n")
        lines.append(content)
        return "\n".join(lines)
    except Exception as e:
        return "[Error] Cannot read template {}: {}".format(p, e)


def _read_docx_template(p: Path, base_dir: Path) -> str:
    """Extract text and placeholders from a .docx template."""
    _try_imports(base_dir / ".." / "vendor")
    if _docx is None:
        return (
            "[Warning] python-docx library not available.\n"
            "Please copy python-docx to vendor/ directory to use .docx templates.\n"
            "For now, treating as binary — cannot extract placeholders."
        )

    try:
        doc = _docx.Document(str(p))
        all_text = []
        for para in doc.paragraphs:
            all_text.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    all_text.append(cell.text)

        combined = "\n".join(all_text)
        placeholders = _extract_placeholders(combined)

        lines = [
            "[Template] {}".format(p.name),
            "Format: .docx (Word)",
            "Paragraphs: {} — Tables: {}".format(len(doc.paragraphs), len(doc.tables)),
        ]
        if placeholders:
            lines.append("Placeholders: {}".format(", ".join("{{{{{}}}}}".format(k) for k in placeholders)))
        lines.append("\n--- Extracted Text ---\n")
        lines.append(combined[:5000])
        return "\n".join(lines)
    except Exception as e:
        return "[Error] Cannot read .docx template {}: {}".format(p, e)


def _read_xlsx_template(p: Path, base_dir: Path) -> str:
    """Extract text and placeholders from an .xlsx template."""
    _try_imports(base_dir / ".." / "vendor")
    if _openpyxl is None:
        return (
            "[Warning] openpyxl library not available.\n"
            "Please copy openpyxl to vendor/ directory to use .xlsx templates.\n"
            "For now, treating as binary — cannot extract placeholders."
        )

    try:
        wb = _openpyxl.load_workbook(str(p), data_only=True)
        all_text = []
        sheet_count = len(wb.sheetnames)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            all_text.append("[Sheet: {}]".format(sheet_name))
            for row in ws.iter_rows(values_only=True):
                row_text = " | ".join(str(v) if v is not None else "" for v in row)
                if row_text.strip():
                    all_text.append(row_text)

        combined = "\n".join(all_text)
        placeholders = _extract_placeholders(combined)

        lines = [
            "[Template] {}".format(p.name),
            "Format: .xlsx (Excel)",
            "Sheets: {}".format(sheet_count),
        ]
        if placeholders:
            lines.append("Placeholders: {}".format(", ".join("{{{{{}}}}}".format(k) for k in placeholders)))
        lines.append("\n--- Extracted Text ---\n")
        lines.append(combined[:5000])
        return "\n".join(lines)
    except Exception as e:
        return "[Error] Cannot read .xlsx template {}: {}".format(p, e)


def _read_pptx_template(p: Path, base_dir: Path) -> str:
    """Extract text and placeholders from a .pptx template."""
    _try_imports(base_dir / ".." / "vendor")
    if _pptx is None:
        return (
            "[Warning] python-pptx library not available.\n"
            "Please copy python-pptx to vendor/ directory to use .pptx templates.\n"
            "For now, treating as binary — cannot extract placeholders."
        )

    try:
        prs = _pptx.Presentation(str(p))
        all_text = []
        slide_count = len(prs.slides)
        for i, slide in enumerate(prs.slides):
            all_text.append("[Slide {}]".format(i + 1))
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = para.text.strip()
                        if t:
                            all_text.append(t)

        combined = "\n".join(all_text)
        placeholders = _extract_placeholders(combined)

        lines = [
            "[Template] {}".format(p.name),
            "Format: .pptx (PowerPoint)",
            "Slides: {}".format(slide_count),
        ]
        if placeholders:
            lines.append("Placeholders: {}".format(", ".join("{{{{{}}}}}".format(k) for k in placeholders)))
        lines.append("\n--- Extracted Text ---\n")
        lines.append(combined[:5000])
        return "\n".join(lines)
    except Exception as e:
        return "[Error] Cannot read .pptx template {}: {}".format(p, e)


# ---------------------------------------------------------------------------
# write_output — extended for .sql and other text formats
# ---------------------------------------------------------------------------

def write_output(path: str, content: str, base_dir: Path) -> str:
    """Write generated content to the output/ directory.

    Creates parent directories as needed. Supports all text formats.
    """
    p = Path(path)
    if not p.is_absolute():
        p = base_dir / "output" / p

    try:
        _ensure_output_dir(p)
        p.write_text(content, encoding="utf-8")
        return "[OK] Output written to {} ({} chars)".format(p, len(content))
    except Exception as e:
        return "[Error] Cannot write output {}: {}".format(p, e)


# ---------------------------------------------------------------------------
# write_docx — Word template filling
# ---------------------------------------------------------------------------

def write_docx(path: str, fields, template_path: str, base_dir: Path) -> str:
    """Fill a Word (.docx) template with field values and save to output.

    Args:
        path: Output .docx file path (relative to output/ or absolute).
        fields: JSON string or dict of key-value pairs for {{KEY}} replacement.
        template_path: Path to the .docx template file.
        base_dir: Base directory for resolving relative paths.

    Returns:
        Result string with status and output path.
    """
    _try_imports(base_dir / ".." / "vendor")
    if _docx is None:
        return (
            "[Error] python-docx library not available.\n"
            "To generate .docx files, copy python-docx to the vendor/ directory.\n"
            "pip install python-docx --target=vendor"
        )

    # Resolve template path
    tp = Path(template_path)
    if not tp.is_absolute():
        tp = base_dir / "templates" / tp

    if not tp.exists():
        return "[Error] Template not found: {}".format(tp)

    # Parse fields
    if isinstance(fields, str):
        try:
            fields = __import__("json").loads(fields)
        except Exception:
            return "[Error] Invalid fields JSON: {}".format(fields[:200])
    if not isinstance(fields, dict):
        return "[Error] Fields must be a JSON object (key-value pairs)."

    # Resolve output path
    op = Path(path)
    if not op.is_absolute():
        op = base_dir / "output" / op

    try:
        _ensure_output_dir(op)

        doc = _docx.Document(str(tp))

        # Replace in paragraphs
        for para in doc.paragraphs:
            for run in para.runs:
                run.text = _replace_placeholders(run.text, fields)

        # Replace in tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.text = _replace_placeholders(run.text, fields)

        doc.save(str(op))
        return "[OK] Document written to {} ({:.1f} KB)".format(
            op, op.stat().st_size / 1024 if op.exists() else 0
        )
    except Exception as e:
        return "[Error] Cannot generate .docx: {}".format(e)


# ---------------------------------------------------------------------------
# write_xlsx — Excel template filling
# ---------------------------------------------------------------------------

def write_xlsx(path: str, fields, template_path: str, base_dir: Path) -> str:
    """Fill an Excel (.xlsx) template with field values and save to output.

    Args:
        path: Output .xlsx file path.
        fields: JSON string or dict of key-value pairs for {{KEY}} replacement.
        template_path: Path to the .xlsx template file.
        base_dir: Base directory for relative paths.

    Returns:
        Result string with status and output path.
    """
    _try_imports(base_dir / ".." / "vendor")
    if _openpyxl is None:
        return (
            "[Error] openpyxl library not available.\n"
            "To generate .xlsx files, copy openpyxl to the vendor/ directory.\n"
            "pip install openpyxl --target=vendor"
        )

    # Resolve template path
    tp = Path(template_path)
    if not tp.is_absolute():
        tp = base_dir / "templates" / tp

    if not tp.exists():
        return "[Error] Template not found: {}".format(tp)

    # Parse fields
    if isinstance(fields, str):
        try:
            fields = __import__("json").loads(fields)
        except Exception:
            return "[Error] Invalid fields JSON: {}".format(fields[:200])
    if not isinstance(fields, dict):
        return "[Error] Fields must be a JSON object (key-value pairs)."

    # Resolve output path
    op = Path(path)
    if not op.is_absolute():
        op = base_dir / "output" / op

    try:
        _ensure_output_dir(op)

        wb = _openpyxl.load_workbook(str(tp))

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        cell.value = _replace_placeholders(cell.value, fields)

        wb.save(str(op))
        return "[OK] Spreadsheet written to {} ({:.1f} KB)".format(
            op, op.stat().st_size / 1024 if op.exists() else 0
        )
    except Exception as e:
        return "[Error] Cannot generate .xlsx: {}".format(e)


# ---------------------------------------------------------------------------
# write_pptx — PowerPoint template filling
# ---------------------------------------------------------------------------

def write_pptx(path: str, fields, template_path: str, base_dir: Path) -> str:
    """Fill a PowerPoint (.pptx) template with field values and save to output.

    Args:
        path: Output .pptx file path.
        fields: JSON string or dict of key-value pairs for {{KEY}} replacement.
        template_path: Path to the .pptx template file.
        base_dir: Base directory for relative paths.

    Returns:
        Result string with status and output path.
    """
    _try_imports(base_dir / ".." / "vendor")
    if _pptx is None:
        return (
            "[Error] python-pptx library not available.\n"
            "To generate .pptx files, copy python-pptx to the vendor/ directory.\n"
            "pip install python-pptx --target=vendor"
        )

    # Resolve template path
    tp = Path(template_path)
    if not tp.is_absolute():
        tp = base_dir / "templates" / tp

    if not tp.exists():
        return "[Error] Template not found: {}".format(tp)

    # Parse fields
    if isinstance(fields, str):
        try:
            fields = __import__("json").loads(fields)
        except Exception:
            return "[Error] Invalid fields JSON: {}".format(fields[:200])
    if not isinstance(fields, dict):
        return "[Error] Fields must be a JSON object (key-value pairs)."

    # Resolve output path
    op = Path(path)
    if not op.is_absolute():
        op = base_dir / "output" / op

    try:
        _ensure_output_dir(op)

        prs = _pptx.Presentation(str(tp))

        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            run.text = _replace_placeholders(run.text, fields)

        prs.save(str(op))
        return "[OK] Presentation written to {} ({:.1f} KB)".format(
            op, op.stat().st_size / 1024 if op.exists() else 0
        )
    except Exception as e:
        return "[Error] Cannot generate .pptx: {}".format(e)
